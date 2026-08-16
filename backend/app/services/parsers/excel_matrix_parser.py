"""
Configurable Multi-Sheet Excel Compliance Matrix Parser.
"""

import openpyxl
import logging
from typing import Dict, Any, List, Optional

logger = logging.getLogger(__name__)


class ConfigurableExcelParser:
    """Dynamic Excel parser supporting custom column maps, header offsets, and multi-sheet iteration."""

    def __init__(
        self,
        file_path: str,
        framework_name: str,
        framework_version: str = "1.0",
        id_col: str = "Control ID",
        title_col: str = "Control Title",
        text_col: str = "Control Specification",
        header_row_offset: int = 1,
        sheets: Optional[List[str]] = None,
    ):
        self.file_path = file_path
        self.framework_name = framework_name
        self.framework_version = framework_version
        self.id_col = id_col.lower().strip()
        self.title_col = title_col.lower().strip()
        self.text_col = text_col.lower().strip()
        self.header_row_offset = header_row_offset
        self.target_sheets = sheets or ["*"]

    def parse(self) -> Dict[str, List[Dict[str, Any]]]:
        wb = openpyxl.load_workbook(self.file_path, data_only=True)
        nodes = []

        sheet_names = wb.sheetnames if "*" in self.target_sheets else [s for s in self.target_sheets if s in wb.sheetnames]

        for sheet_name in sheet_names:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) < self.header_row_offset:
                continue

            # Identify column headers at specified offset
            header_row = [str(c).lower().strip() if c is not None else "" for c in rows[self.header_row_offset - 1]]
            
            id_idx = next((i for i, h in enumerate(header_row) if self.id_col in h), 0)
            title_idx = next((i for i, h in enumerate(header_row) if self.title_col in h), 1)
            text_idx = next((i for i, h in enumerate(header_row) if self.text_col in h), 2)

            for row in rows[self.header_row_offset:]:
                if not row or len(row) <= id_idx or not row[id_idx]:
                    continue
                ctrl_id = str(row[id_idx]).strip()
                # Skip if header repeated or empty
                if ctrl_id.lower() == self.id_col or not ctrl_id:
                    continue

                ctrl_title = str(row[title_idx]).strip() if len(row) > title_idx and row[title_idx] else ctrl_id
                ctrl_text = str(row[text_idx]).strip() if len(row) > text_idx and row[text_idx] else ""

                nodes.append({
                    "framework_obj_id": f"{self.framework_name}:{ctrl_id}",
                    "framework_name": self.framework_name,
                    "framework_version": self.framework_version,
                    "objective_name": ctrl_title,
                    "objective_text": ctrl_text,
                    "sheet_category": sheet_name,
                })

        logger.info("ConfigurableExcelParser extracted %d nodes from %s", len(nodes), self.file_path)
        return {"nodes": nodes, "edges": []}
